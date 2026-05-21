from fastapi import FastAPI
from fastapi.responses import FileResponse
from docx import Document
import uuid
import os

app = FastAPI()

# 自動建立 downloads 資料夾
os.makedirs("downloads", exist_ok=True)

@app.get("/generate-doc")
def generate_doc():

    # 隨機檔名
    filename = f"{uuid.uuid4()}.docx"

    # 檔案路徑
    path = f"downloads/{filename}"

    # 建立 Word 文件
    doc = Document()

    doc.add_heading("生理學題目", level=1)

    doc.add_paragraph("1. 人體最大的器官是皮膚")
    doc.add_paragraph("2. 心臟有四個腔室")

    # 儲存 Word
    doc.save(path)

    # 回傳下載網址
    return {
        "message": "檔案已生成",
        "download_url": f"https://maiagent-download.onrender.com/download/{filename}"
    }

# 下載 API
@app.get("/download/{filename}")
def download(filename: str):

    path = f"downloads/{filename}"

    return FileResponse(
        path,
        filename=filename,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )